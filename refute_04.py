import openpyxl, datetime
from collections import defaultdict

P = "C:/Users/usuario/Downloads/Registro KV Excel.xlsx"
wbv = openpyxl.load_workbook(P, data_only=True)

def mk(v):
    if isinstance(v, (datetime.datetime, datetime.date)): return (v.year, v.month)
    return None

def rows(ws, r1, r2):
    w = wbv[ws]; out=[]
    for r in range(r1, r2+1):
        g = w.cell(r,7).value
        if not isinstance(g,(int,float)): continue
        out.append(dict(row=r, fecha=w.cell(r,1).value, cod=w.cell(r,2).value,
                        desc=w.cell(r,3).value, talla=w.cell(r,4).value,
                        e=w.cell(r,5).value, f=w.cell(r,6).value, g=g))
    return out

S1=rows("STOCK 1",197,454); R1=rows("Registro de STOCK 1",197,418)
S2=rows("STOCK 2",600,695); R2=rows("Registro de STOCK 2",600,695)
def keyof(d): return (str(d["cod"]).strip() if d["cod"] is not None else None,
                      str(d["desc"]).strip() if d["desc"] is not None else None,
                      str(d["talla"]).strip() if d["talla"] is not None else None)

print("="*72)
print("1) EVERY row whose G changed: is ANY of them dated in AUGUST?")
tot_pre=0.0; tot_aug=0.0; allch=[]
for nm,CUR,PREV in (("S1",S1,R1),("S2",S2,R2)):
    pm=defaultdict(list)
    for d in PREV: pm[keyof(d)].append(d)
    for d in CUR:
        k=keyof(d)
        if pm.get(k):
            p=pm[k].pop(0)
            if abs(p["g"]-d["g"])>0.004:
                dl=d["g"]-p["g"]; allch.append((nm,d,p,dl))
                if mk(d["fecha"])==(2026,8): tot_aug+=dl
                else: tot_pre+=dl
print(f"   changed rows: {len(allch)}   net delta = {round(tot_pre+tot_aug,2):,.2f}")
print(f"   delta on rows dated BEFORE august = {round(tot_pre,2):,.2f}")
print(f"   delta on rows dated IN     august = {round(tot_aug,2):,.2f}")
print(f"   -> analyst calls 17,049.28 'crecimiento fechado antes de agosto'")
print(f"      actual growth from pre-august-dated rows = {round(tot_pre,2):,.2f}   DIFF = {round(tot_pre-17049.28,2):,.2f}")

print()
print("="*72)
print("2) The 4,537.38 august-dated rows: were they in the july photo? (3,814.80 + 722.58 claim)")
for nm,CUR,PREV in (("STOCK 1",S1,R1),("STOCK 2",S2,R2)):
    pm=defaultdict(list)
    for d in PREV: pm[keyof(d)].append(d)
    was=0.0; wasn=0; new=0.0; newn=0; prevval=0.0
    for d in CUR:
        if mk(d["fecha"])!=(2026,8): continue
        k=keyof(d)
        if pm.get(k):
            p=pm[k].pop(0); was+=d["g"]; wasn+=1; prevval+=p["g"]
        else:
            new+=d["g"]; newn+=1
    print(f"   {nm}: aug-dated rows already in july photo: n={wasn} Gnow={was:,.2f} (Gthen={prevval:,.2f}); NOT in photo: n={newn} G={new:,.2f}")
    globals()[nm[-1]+"_was"]=was; globals()[nm[-1]+"_new"]=new; globals()[nm[-1]+"_then"]=prevval
print(f"   TOTAL aug-dated already in july photo = {round(globals()['1_was']+globals()['2_was'],2):,.2f}")
print(f"   TOTAL aug-dated NOT in july photo     = {round(globals()['1_new']+globals()['2_new'],2):,.2f}")
print(f"   -> analyst claims split 3,814.80 (in july photo) + 722.58 (costs added in august)")

print()
print("="*72)
print("3) Column F (COSTOS) on the august-dated rows")
cf1=round(sum((d["f"] or 0) for d in S1 if mk(d["fecha"])==(2026,8) and isinstance(d["f"],(int,float))),2)
cf2=round(sum((d["f"] or 0) for d in S2 if mk(d["fecha"])==(2026,8) and isinstance(d["f"],(int,float))),2)
ce1=round(sum((d["e"] or 0) for d in S1 if mk(d["fecha"])==(2026,8) and isinstance(d["e"],(int,float))),2)
ce2=round(sum((d["e"] or 0) for d in S2 if mk(d["fecha"])==(2026,8) and isinstance(d["e"],(int,float))),2)
print(f"   aug-dated  E(precio compra) = {ce1:,.2f} + {ce2:,.2f} = {round(ce1+ce2,2):,.2f}")
print(f"   aug-dated  F(costos)        = {cf1:,.2f} + {cf2:,.2f} = {round(cf1+cf2,2):,.2f}")
print(f"   aug-dated  E+F              = {round(ce1+ce2+cf1+cf2,2):,.2f}   (G total 4,537.38)")
print(f"   -> is F = 722.58 ?  {'YES' if abs(round(cf1+cf2,2)-722.58)<0.005 else 'NO, it is %.2f'%(cf1+cf2)}")

print()
print("="*72)
print("4) FLUJO sheet: august COMPRA rows in soles (the other 3,814.80)")
w=wbv["FLUJO"]
agg=defaultdict(lambda:[0.0,0.0,0.0,0])
for r in range(2, w.max_row+1):
    f=w.cell(r,1).value; acc=w.cell(r,2).value
    if mk(f)!=(2026,8) or not acc: continue
    ing=w.cell(r,6).value or 0; egr=w.cell(r,7).value or 0; egu=w.cell(r,8).value or 0
    a=agg[str(acc).strip().upper()]
    a[0]+= ing if isinstance(ing,(int,float)) else 0
    a[1]+= egr if isinstance(egr,(int,float)) else 0
    a[2]+= egu if isinstance(egu,(int,float)) else 0
    a[3]+=1
for k in sorted(agg):
    print(f"   {k:<12} n={agg[k][3]:<4} ingreso={agg[k][0]:>11,.2f}  egreso S/={agg[k][1]:>10,.2f}  egreso US$={agg[k][2]:>9,.2f}")

print()
print("="*72)
print("5) THE ALGEBRA behind '3,814.80 rebaja de las constantes'")
a=round(sum(d["g"] for d in S1),2); b=round(sum(d["g"] for d in R1),2)
c=round(sum(d["g"] for d in S2),2); dd=round(sum(d["g"] for d in R2),2)
print(f"   a=SUM S1 now {a:,.2f}   b=SUM RegS1 {b:,.2f}   c=SUM S2 now {c:,.2f}   d=SUM RegS2 {dd:,.2f}")
print(f"   21,586.66 = (a-84,285.52)+(c-234)")
print(f"   17,771.86 = (a-b)+(c-d)")
print(f"   => residual = (b+d) - (84,285.52+234) = {b:,.2f}+{dd:,.2f} - 84,519.52 = {round(b+dd-84519.52,2):,.2f}")
print(f"   So 3,814.80 is a RESIDUAL between hardcoded constants and the real prior sums,")
print(f"   not a 'rebaja de constantes'. Actual constant movement: 72,135.87 -> 84,519.52 = {round(84519.52-72135.87,2):+,.2f}")
