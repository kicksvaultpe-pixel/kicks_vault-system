import openpyxl, datetime
from collections import defaultdict

P = "C:/Users/usuario/Downloads/Registro KV Excel.xlsx"
wbv = openpyxl.load_workbook(P, data_only=True)
def mk(v):
    if isinstance(v,(datetime.datetime,datetime.date)): return (v.year,v.month)
    return None

def bymonth(ws,r1,r2):
    w=wbv[ws]; agg=defaultdict(lambda:[0.0,0]); tot=0.0
    for r in range(r1,r2+1):
        g=w.cell(r,7).value
        if not isinstance(g,(int,float)): continue
        k=mk(w.cell(r,1).value); agg[k][0]+=g; agg[k][1]+=1; tot+=g
    return agg, round(tot,2)

A1,a = bymonth("STOCK 1",197,454)
B1,b = bymonth("Registro de STOCK 1",197,418)
A2,c = bymonth("STOCK 2",600,695)
B2,d = bymonth("Registro de STOCK 2",600,695)

print("="*74)
print("MATCHING-FREE VERIFICATION (pure date grouping, no row pairing)")
aug=(2026,8)
s1n=round(A1[aug][0],2); s1t=round(B1[aug][0],2)
s2n=round(A2[aug][0],2); s2t=round(B2[aug][0],2)
print(f"  AUG-dated G, NOW   : S1={s1n:,.2f} (n={A1[aug][1]})  S2={s2n:,.2f} (n={A2[aug][1]})  tot={round(s1n+s2n,2):,.2f}")
print(f"  AUG-dated G, JULY  : S1={s1t:,.2f} (n={B1[aug][1]})  S2={s2t:,.2f} (n={B2[aug][1]})  tot={round(s1t+s2t,2):,.2f}")
print(f"  -> growth ON aug-dated rows          = {round(s1n+s2n-s1t-s2t,2):,.2f}")
pre_now = round(a-s1n + c-s2n, 2)
pre_then= round(b-s1t + d-s2t, 2)
print(f"  PRE-aug-dated G, NOW  = {pre_now:,.2f}")
print(f"  PRE-aug-dated G, JULY = {pre_then:,.2f}")
print(f"  -> growth on PRE-aug rows            = {round(pre_now-pre_then,2):,.2f}   [analyst: 17,049.28]")
print()
print(f"  CHECK constants: b - 84,285.52 = {round(b-84285.52,2):,.2f}  (should equal S1 aug-dated in july photo {s1t:,.2f})")
print(f"  CHECK constants: d - 234.00    = {round(d-234,2):,.2f}  (should equal S2 aug-dated in july photo {s2t:,.2f})")
print(f"  => the two constants are EXACTLY 'prior sum minus the august-dated rows already sitting in the sheet'")

print()
print("="*74)
print("ARE THE TWO 'CAMINOS' INDEPENDENT?  (algebraic relation)")
G=round(a-b+c-d,2); Gpre=round(pre_now-pre_then,2); Gaug=round(s1n+s2n-s1t-s2t,2)
K=round((b-84285.52)+(d-234),2); AUG=round(s1n+s2n,2)
print(f"  Camino 1:  G(total growth) {G:,.2f} + K(constant gap) {K:,.2f} = {round(G+K,2):,.2f}")
print(f"  Camino 2:  Gpre {Gpre:,.2f} + AUGnow {AUG:,.2f} = {round(Gpre+AUG,2):,.2f}")
print(f"  but  G = Gpre + Gaug  ({Gpre:,.2f} + {Gaug:,.2f} = {G:,.2f})")
print(f"  and  AUGnow = K + Gaug ({K:,.2f} + {Gaug:,.2f} = {AUG:,.2f})")
print(f"  => Camino 2 is Camino 1 with the SAME {Gaug:,.2f} term moved across. One identity, not two.")

print()
print("="*74)
print("HOW OLD IS THE 17,049.28? (growth by the month the row is DATED)")
allm=set(list(A1)+list(B1)+list(A2)+list(B2))
tot=0.0
for k in sorted([x for x in allm if x is not None]):
    dl = (A1[k][0]-B1[k][0]) + (A2[k][0]-B2[k][0])
    if abs(dl)>0.004:
        print(f"    rows dated {k}: growth {dl:>11,.2f}")
        tot+=dl
dln=(A1[None][0]-B1[None][0])+(A2[None][0]-B2[None][0])
if abs(dln)>0.004: print(f"    rows dated (blank) : growth {dln:>11,.2f}")
print(f"    TOTAL {round(tot+dln,2):,.2f}")

print()
print("="*74)
print("BROKEN CROSS-SHEET REFERENCE IN THE JULY SNAPSHOT")
print(f"  'Registro de STOCK 1'!O420 = SUM('STOCK 2'!G600:G695)  -> points at the LIVE sheet")
print(f"     live  STOCK 2 sum today          = {c:,.2f}   (cached in O420 = 4,501.57)")
print(f"     frozen 'Registro de STOCK 2' sum = {d:,.2f}   <- the real july value")
print(f"  => the july snapshot total 20,085.89 is CONTAMINATED by august data.")
print(f"     genuine july total = 15,584.32 + {d:,.2f} = {round(15584.32+d,2):,.2f}")
print(f"     overstated by {round(c-d,2):,.2f}")
