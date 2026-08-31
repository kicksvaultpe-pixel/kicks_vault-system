# -*- coding: utf-8 -*-
import openpyxl, datetime
p = "C:/Users/usuario/Downloads/Registro KV Excel.xlsx"
wv = openpyxl.load_workbook(p, data_only=True)

def num(x):
    return float(x) if isinstance(x, (int, float)) else 0.0
def mes(d):
    if isinstance(d, datetime.datetime): return (d.year, d.month)
    return None

print("### DESCOMPOSICION EXACTA ###")
s1_now, s1_jul = 101604.61, 87720.19
s2_now, s2_jul = 4501.57, 614.13
g1 = s1_now - s1_jul; g2 = s2_now - s2_jul
print(f"  crecimiento real S1 = {s1_now:,.2f} - {s1_jul:,.2f} = {g1:,.2f}")
print(f"  crecimiento real S2 = {s2_now:,.2f} - {s2_jul:,.2f} = {g2:,.2f}")
print(f"  crecimiento total   = {g1+g2:,.2f}")
print(f"  rebaja S1 = {s1_jul-84285.52:,.2f} ; rebaja S2 = {s2_jul-234.0:,.2f} ; total = {(s1_jul-84285.52)+(s2_jul-234.0):,.2f}")
print(f"  COMPRAS MES = {g1+g2:,.2f} + {(s1_jul-84285.52)+(s2_jul-234.0):,.2f} = {g1+g2+(s1_jul-84285.52)+(s2_jul-234.0):,.2f}  (Excel O457 = 21,586.66)")

print()
print("### FLUJO: todas las filas de AGOSTO 2026 ###")
ws = wv["FLUJO"]
tot_sol = {}; tot_usd = {}; nrows = {}
ago_compra_soles = []
for r in range(2, ws.max_row+1):
    a = ws[f"A{r}"].value
    if mes(a) != (2026, 8): continue
    acc = (ws[f"B{r}"].value or "").strip() if isinstance(ws[f"B{r}"].value, str) else ws[f"B{r}"].value
    ing, egr, egrusd = num(ws[f"F{r}"].value), num(ws[f"G{r}"].value), num(ws[f"H{r}"].value)
    nrows[acc] = nrows.get(acc,0)+1
    tot_sol[acc] = tot_sol.get(acc,0)+egr
    tot_usd[acc] = tot_usd.get(acc,0)+egrusd
    if acc == "COMPRA" and egr:
        ago_compra_soles.append((r, a.date(), ws[f"C{r}"].value, ws[f"D{r}"].value, egr, egrusd, ws[f"I{r}"].value))
for k in sorted(nrows, key=lambda x: str(x)):
    print(f"  {str(k):10s} n={nrows[k]:3d}  egreso_soles={tot_sol[k]:12,.2f}  egreso_usd={tot_usd[k]:10,.2f}")

print()
print("### Filas COMPRA de agosto CON importe en soles ###")
s = 0.0
for r, d, cod, desc, egr, egrusd, info in ago_compra_soles:
    s += egr
    print(f"  fila {r:4d} {d} {str(cod):10s} S/{egr:10,.2f}  US${egrusd:8,.2f}  info={info!r}  desc={str(desc)[:40]!r}")
print(f"  SUMA = {s:,.2f}   (claim: 3,814.80)")
