# -*- coding: utf-8 -*-
import openpyxl, datetime
p = "C:/Users/usuario/Downloads/Registro KV Excel.xlsx"
wv = openpyxl.load_workbook(p, data_only=True)

def num(x):
    return x if isinstance(x, (int, float)) else 0.0

def dump(sheet, r0, r1, label):
    ws = wv[sheet]
    tot = 0.0
    rows = []
    for r in range(r0, r1+1):
        a = ws[f"A{r}"].value
        b = ws[f"B{r}"].value
        g = num(ws[f"G{r}"].value)
        e = num(ws[f"E{r}"].value)
        f_ = num(ws[f"F{r}"].value)
        h = ws[f"H{r}"].value
        tot += g
        rows.append((r, a, b, e, f_, g, h))
    print(f"{label}: {sheet}!G{r0}:G{r1}  SUM = {tot:,.2f}   ({len(rows)} filas)")
    return rows, tot

print("="*70)
r_s1_now, t_s1_now = dump("STOCK 1", 197, 454, "ACTUAL   S1")
r_s1_jul, t_s1_jul = dump("Registro de STOCK 1", 197, 418, "JULIO    S1")
r_s1_now418, t_s1_now418 = dump("STOCK 1", 197, 418, "ACTUAL S1 (solo 197:418)")
print()
r_s2_now, t_s2_now = dump("STOCK 2", 600, 695, "ACTUAL   S2")
r_s2_jul, t_s2_jul = dump("Registro de STOCK 2", 600, 695, "JULIO    S2")
print()
print(f"Constante S1 agosto = 84285.52 ; SUM julio G197:G418 = {t_s1_jul:,.2f}")
print(f"  diferencia (rebaja S1) = {t_s1_jul - 84285.52:,.2f}")
print(f"Constante S2 agosto = 234.00   ; SUM julio S2 = {t_s2_jul:,.2f} ; SUM actual S2 = {t_s2_now:,.2f}")
print(f"  crecimiento real rango S2 = {t_s2_now - t_s2_jul:,.2f}")
print()
print("--- extension real de cada hoja ---")
for sh in ("STOCK 1","STOCK 2","Registro de STOCK 1","Registro de STOCK 2"):
    ws = wv[sh]
    last = 0
    for r in range(1, ws.max_row+1):
        if ws[f"B{r}"].value or ws[f"G{r}"].value:
            last = r
    print(f"  {sh}: max_row={ws.max_row}, ultima fila con B o G = {last}")
