# -*- coding: utf-8 -*-
import openpyxl
P = "C:/Users/usuario/Downloads/Registro KV Excel.xlsx"
wbF = openpyxl.load_workbook(P, data_only=False)
wbV = openpyxl.load_workbook(P, data_only=True)

def show(sheet, rows, cols=(1,2,3,5,6,7,8)):
    f=wbF[sheet]; v=wbV[sheet]
    for r in rows:
        line=[]
        for c in cols:
            fv=f.cell(r,c).value; vv=v.cell(r,c).value
            line.append("%s{F:%r|V:%r}"%(openpyxl.utils.get_column_letter(c),fv,vv))
        print(sheet, r, " ".join(line))

print("=== STOCK 1 filas clave ===")
show("STOCK 1",[364,394])
print("=== Registro de STOCK 1 filas clave ===")
show("Registro de STOCK 1",[364,394])
print("=== STOCK 2 E600 ===")
show("STOCK 2",[600])
print("=== Registro de STOCK 2 E600 ===")
show("Registro de STOCK 2",[600])

print("\n=== COMPRA fila 475 y alrededores ===")
cf=wbF["COMPRA"]; cv=wbV["COMPRA"]
for r in [1,474,475,476]:
    print(r,[cv.cell(r,c).value for c in range(1,9)])
    print("   F:",[cf.cell(r,c).value for c in range(1,9)])
