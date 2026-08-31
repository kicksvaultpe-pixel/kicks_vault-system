# -*- coding: utf-8 -*-
import openpyxl, datetime
from collections import Counter

P = "C:/Users/usuario/Downloads/Registro KV Excel.xlsx"
wbv = openpyxl.load_workbook(P, data_only=True)

def num(x):
    return x if isinstance(x, (int, float)) and not isinstance(x, bool) else None

def analyze(sheet, r1, r2, label):
    ws = wbv[sheet]
    print(f"\n===== {label}: {sheet}!G{r1}:G{r2}  (max_row={ws.max_row}) =====")
    tot = 0.0; n = 0
    months = Counter()
    rows = []
    for r in range(r1, r2 + 1):
        g = num(ws.cell(row=r, column=7).value)
        a = ws.cell(row=r, column=1).value
        b = ws.cell(row=r, column=2).value
        c = ws.cell(row=r, column=3).value
        h = ws.cell(row=r, column=8).value
        if g is not None:
            tot += g; n += 1
            m = a.strftime("%Y-%m") if isinstance(a, datetime.datetime) else f"NOFECHA({a!r})"
            months[m] += 1
            rows.append((r, m, b, str(c)[:32], g, h))
    print(f"  SUM = {tot:.2f}   celdas numericas = {n}")
    print(f"  por mes: {sorted(months.items())}")
    return tot, n, rows, ws

# --- STOCK 2 live y snapshot ---
t_live, n_live, rows_live, ws_live = analyze('STOCK 2', 600, 695, "LIVE (agosto)")
t_snap, n_snap, rows_snap, ws_snap = analyze('Registro de STOCK 2', 600, 695, "SNAPSHOT (julio)")

print(f"\n  DIFERENCIA live - snapshot = {t_live - t_snap:.2f}")
print(f"  4267.57 + 234 = {4267.57+234:.2f}   (=SUM live?) {abs(t_live-4501.57)<0.01}")
print(f"  Ajuste propuesto por el analista: 614.13 - 234 = {614.13-234:.2f}")

# --- extension de STOCK 2 mas alla de 695 ---
print("\n===== STOCK 2: filas 696+ con dato =====")
for r in range(696, ws_live.max_row + 1):
    g = num(ws_live.cell(row=r, column=7).value)
    a = ws_live.cell(row=r, column=1).value
    if g is not None:
        print(f"  fila {r}: A={a} G={g} B={ws_live.cell(row=r,column=2).value}")

print("\n===== STOCK 2: primeras filas antes de 600 con dato (ultimas 8) =====")
tmp = []
for r in range(2, 600):
    g = num(ws_live.cell(row=r, column=7).value)
    if g is not None:
        tmp.append(r)
print(f"  filas con G no vacio antes de 600: {len(tmp)}, ultimas: {tmp[-8:] if tmp else None}")
