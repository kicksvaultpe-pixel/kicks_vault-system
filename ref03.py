# -*- coding: utf-8 -*-
import openpyxl
from decimal import Decimal

P = "C:/Users/usuario/Downloads/Registro KV Excel.xlsx"
wbv = openpyxl.load_workbook(P, data_only=True)

S = wbv["STOCK 1"]; R = wbv["Registro de STOCK 1"]

def row(ws, r):
    return tuple(ws.cell(row=r, column=c).value for c in range(1,9))

print("=== COMPARACION FILA A FILA  STOCK 1  vs  Registro de STOCK 1 (197:418) ===")
difs = 0
delta_G = Decimal(0)
for r in range(197, 419):
    a = row(S, r); b = row(R, r)
    if a != b:
        difs += 1
        ga = a[6] if isinstance(a[6], (int,float)) else 0
        gb = b[6] if isinstance(b[6], (int,float)) else 0
        delta_G += Decimal(str(ga)) - Decimal(str(gb))
        if difs <= 40:
            print("  fila %d" % r)
            print("     STOCK1: %r" % (a,))
            print("     REG   : %r" % (b,))
print("filas distintas: %d ; delta en G = %s" % (difs, delta_G))

def suma(ws, r1, r2):
    s = Decimal(0); n=0
    for r in range(r1,r2+1):
        v = ws.cell(row=r, column=7).value
        if isinstance(v,(int,float)): s += Decimal(str(v)); n+=1
    return s,n

a1,na = suma(S,197,418); b1,nb = suma(R,197,418)
a2,na2 = suma(S,419,454)
print()
print("SUM STOCK1 G197:G418 = %s (%d)" % (a1,na))
print("SUM REG    G197:G418 = %s (%d)" % (b1,nb))
print("SUM STOCK1 G419:G454 = %s (%d)  <- filas nuevas de agosto" % (a2,na2))
print("a1+a2 = %s" % (a1+a2))
print()
print("Crecimiento real del rango vs foto julio: 101604.61 - 87720.19 = %s" % (Decimal("101604.61")-Decimal("87720.19")))
print("Lo que dice la formula:                    101604.61 - 84285.52 = %s" % (Decimal("101604.61")-Decimal("84285.52")))
