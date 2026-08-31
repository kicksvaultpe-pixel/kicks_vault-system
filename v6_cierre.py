# -*- coding: utf-8 -*-
import openpyxl
from datetime import datetime
P = "C:/Users/usuario/Downloads/Registro KV Excel.xlsx"
wbV = openpyxl.load_workbook(P, data_only=True)
def num(x): return float(x) if isinstance(x,(int,float)) else 0.0
def esago(f): return isinstance(f,datetime) and f.year==2026 and f.month==8

print("### 1. De donde salen las constantes 84285.52 y 234 ###")
for live,snap,r0,rL,rS,K in [("STOCK 1","Registro de STOCK 1",197,454,418,84285.52),
                             ("STOCK 2","Registro de STOCK 2",600,695,695,234.0)]:
    for name,sh,r1 in [("LIVE",live,rL),("SNAP",snap,rS)]:
        v=wbV[sh]; tot=0; ago=0
        for r in range(r0,r1+1):
            if not v.cell(r,2).value: continue
            g=num(v.cell(r,7).value); tot+=g
            if esago(v.cell(r,1).value): ago+=g
        print(f"  {sh:22} SUM={tot:10,.2f}  agosto={ago:9,.2f}  NO-agosto={tot-ago:10,.2f}")
    print(f"  --> constante usada = {K:,.2f}")
print()

print("### 2. Cierre exacto del gap ###")
E_ovr_s1=12045.94; E_ovr_s2=3828.14
v1=wbV["STOCK 1"]; v2=wbV["STOCK 2"]
def rng(v,r0,r1):
    t=a=0
    for r in range(r0,r1+1):
        if not v.cell(r,2).value: continue
        g=num(v.cell(r,7).value); t+=g
        if esago(v.cell(r,1).value): a+=g
    return t,a
t1,a1=rng(v1,197,454); t2,a2=rng(v2,600,695)
excel = (t1-84285.52)+(t2-234)
app   = 4157.25+380.13
print(f"  Excel COMPRAS MES  = ({t1:,.2f} - 84,285.52) + ({t2:,.2f} - 234.00) = {excel:,.2f}")
print(f"  App (filas fechadas agosto) = {app:,.2f}")
print(f"  GAP = {excel-app:,.2f}")
ex1=(t1-a1)-84285.52; ex2=(t2-a2)-234
print(f"  exceso STOCK 1 = {ex1:,.2f}   (overrides E {E_ovr_s1:,.2f} + resto {ex1-E_ovr_s1:,.2f})")
print(f"  exceso STOCK 2 = {ex2:,.2f}   (overrides E {E_ovr_s2:,.2f} + resto {ex2-E_ovr_s2:,.2f})")
print(f"  TOTAL overrides E = {E_ovr_s1+E_ovr_s2:,.2f}   TOTAL resto (COSTOS col F) = {(ex1+ex2)-(E_ovr_s1+E_ovr_s2):,.2f}")
print()

print("### 3. Los 20 codigos en la hoja FLUJO ###")
fv=wbV["FLUJO"]
cods=set("P-00964 P-00967 P-00969 P-00972 P-00978 P-00979 P-00980 P-00981 P-00982 P-00983 P-00986 P-00988 P-00996 P-00961 P-00962 P-00963 P-00975 P-00976 P-00984 P-00987".split())
hits={}
for r in range(2,fv.max_row+1):
    c=fv.cell(r,3).value
    if c in cods:
        hits.setdefault(c,[]).append((r,fv.cell(r,1).value,fv.cell(r,2).value,num(fv.cell(r,6).value),num(fv.cell(r,7).value),num(fv.cell(r,8).value),fv.cell(r,9).value))
for c in sorted(cods):
    for h in hits.get(c,[("--","SIN FILA EN FLUJO","","","","","")]):
        print(f"  {c} r{h[0]} fecha={str(h[1])[:10]} accion={h[2]} ing={h[3]} egrS/={h[4]} egr$={h[5]} info={h[6]}")
