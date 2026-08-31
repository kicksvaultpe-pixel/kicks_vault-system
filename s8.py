import openpyxl, datetime
p="C:/Users/usuario/Downloads/Registro KV Excel.xlsx"
wb=openpyxl.load_workbook(p,data_only=True)
def s(v): return v if isinstance(v,(int,float)) else 0.0
targets=set("""P-00961 P-00962 P-00963 P-00964 P-00967 P-00969 P-00972 P-00975 P-00976 P-00978 P-00979
P-00980 P-00981 P-00982 P-00983 P-00984 P-00986 P-00987 P-00988 P-00989 P-00990 P-00991 P-00992 P-00993
P-00996 P-00629 P-00630 P-00633 P-00634 P-00635 P-00636 P-00637 P-00638""".split())
for sh in ['FLUJO','Registro de FLUJO']:
    w=wb[sh]; print("=====",sh,"max_row",w.max_row)
    ds=[]; ago=[]
    for r in range(2,w.max_row+1):
        a=w['A'+str(r)].value
        if isinstance(a,datetime.datetime): ds.append(a)
        b=w['B'+str(r)].value
        if isinstance(a,datetime.datetime) and a.year==2026 and a.month==8:
            ago.append((r,a,b,w['C'+str(r)].value,s(w['F'+str(r)].value),s(w['G'+str(r)].value),s(w['H'+str(r)].value),w['I'+str(r)].value))
    if ds: print("  fechas:",min(ds).strftime('%Y-%m-%d'),"..",max(ds).strftime('%Y-%m-%d'))
    print("  filas agosto:",len(ago))
    tot={}
    for r,a,b,c,f,g,h,i in ago:
        k=str(b); t=tot.setdefault(k,[0,0,0,0]); t[0]+=1; t[1]+=f; t[2]+=g; t[3]+=h
    for k in sorted(tot): print("   %-10s n=%-4d ING=%9.2f EGR_S=%9.2f EGR_USD=%9.2f"%(k,tot[k][0],tot[k][1],tot[k][2],tot[k][3]))
    hit=[x for x in ago if str(x[3]) in targets]
    print("  filas agosto que referencian los CODIGOS con delta pre-agosto:",len(hit))
    for x in hit[:40]: print("     r%d %s %s %s ING=%.2f EGR_S=%.2f EGR_USD=%.2f | %s"%(x[0],x[1].strftime('%Y-%m-%d'),x[2],x[3],x[4],x[5],x[6],str(x[7])[:60]))
