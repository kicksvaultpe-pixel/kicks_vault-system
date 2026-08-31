# -*- coding: utf-8 -*-
import openpyxl
from decimal import Decimal
from collections import defaultdict

P = "C:/Users/usuario/Downloads/Registro KV Excel.xlsx"
wbv = openpyxl.load_workbook(P, data_only=True)
S = wbv["STOCK 1"]; R = wbv["Registro de STOCK 1"]
S2 = wbv["STOCK 2"]; R2 = wbv["Registro de STOCK 2"]

def d(v):
    return Decimal(str(v)) if isinstance(v,(int,float)) else Decimal(0)

print("=== 1) QUE HAY EN LAS FILAS 419:454 de STOCK 1 (la 'extension' del rango) ===")
for r in range(419, 455):
    vals = [S.cell(row=r,column=c).value for c in range(1,9)]
    if any(v not in (None,'',0) for v in vals):
        print("  fila %d: %r" % (r, vals))
print("  (todo lo demas vacio o cero)")

print()
print("=== 2) RANGO DE FECHAS DE CADA HOJA ===")
for nm, ws, r1, r2 in [("STOCK 1", S,197,454), ("Registro de STOCK 1", R,197,418),
                        ("STOCK 2", S2,600,695), ("Registro de STOCK 2", R2,600,695)]:
    ds = [ws.cell(row=r,column=1).value for r in range(r1,r2+1)]
    ds = [x for x in ds if hasattr(x,'year')]
    if ds:
        print("  %-22s  n=%3d  min=%s  max=%s" % (nm, len(ds), min(ds).date(), max(ds).date()))
    # ultima fila con codigo
    last = None
    for r in range(r2, r1-1, -1):
        if ws.cell(row=r,column=2).value: last = r; break
    print("       ultima fila con CODIGO: %s -> %r" % (last, ws.cell(row=last,column=2).value if last else None))

print()
print("=== 3) DELTA DE G (STOCK1 - REG) POR MES DE LA FECHA, filas 197:418 ===")
by = defaultdict(lambda: [Decimal(0),0])
tot = Decimal(0)
augdelta = Decimal(0)
for r in range(197,419):
    dg = d(S.cell(row=r,column=7).value) - d(R.cell(row=r,column=7).value)
    if dg == 0: continue
    f = S.cell(row=r,column=1).value
    k = "%04d-%02d" % (f.year,f.month) if hasattr(f,'year') else "sin fecha"
    by[k][0] += dg; by[k][1] += 1
    tot += dg
    if k.startswith("2026-08"): augdelta += dg
for k in sorted(by):
    print("   %-10s  delta=%10s   (%d filas)" % (k, by[k][0], by[k][1]))
print("   TOTAL delta = %s" % tot)
print("   delta de filas fechadas AGOSTO = %s" % augdelta)
print("   delta de filas fechadas <= JULIO = %s" % (tot - augdelta))

print()
print("=== 4) BUSCANDO EL ORIGEN DE 3434.67 ===")
print("   87720.19 - 84285.52 = %s" % (Decimal("87720.19")-Decimal("84285.52")))
# suma de G en REG para filas fechadas en agosto
regaug = Decimal(0); n=0
for r in range(197,419):
    f = R.cell(row=r,column=1).value
    if hasattr(f,'year') and (f.year,f.month)==(2026,8):
        regaug += d(R.cell(row=r,column=7).value); n+=1
print("   SUM G en 'Registro de STOCK 1' de filas fechadas AGOSTO = %s (%d filas)" % (regaug,n))
print("   87720.19 - %s = %s" % (regaug, Decimal("87720.19")-regaug))
