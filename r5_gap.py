# -*- coding: utf-8 -*-
import openpyxl, datetime, collections
P = "C:/Users/usuario/Downloads/Registro KV Excel.xlsx"
wb = openpyxl.load_workbook(P, data_only=True)

# --- 1. los 11 codigos borrados: existen en algun lado del FLUJO vivo?
def loadF(s):
    ws=wb[s]; rows=[]
    for r in range(2, ws.max_row+1):
        v=[ws.cell(r,c).value for c in range(1,10)]
        if all(x is None for x in v[:8]): continue
        rows.append(dict(row=r,fecha=v[0],accion=v[1],cod=str(v[2]).strip() if v[2] else None,
                         ing=v[5],egr=v[6],usd=v[7]))
    return rows
fl=loadF("FLUJO")
borrados=['P-00965','P-00966','P-00968','P-00970','P-00971','P-00973','P-00974','P-00977','P-00985','P-00994','P-00995']
hit=[r for r in fl if r['cod'] in borrados]
print("### Los 11 codigos jun/jul que en la foto tenian SOLES, buscados en TODO el FLUJO vivo:")
print("   filas encontradas:", len(hit))
for r in hit: print("   ", r)
print()

# --- 2. diff completo de STOCK por codigo
def loadS(s, r0, r1):
    ws=wb[s]; rows=[]
    for r in range(r0, r1+1):
        v=[ws.cell(r,c).value for c in range(1,9)]
        if all(x is None for x in v): continue
        rows.append(dict(sh=s,row=r,fecha=v[0],cod=str(v[1]).strip() if v[1] else None,
                         pc=v[4],costos=v[5],pf=v[6],estado=v[7]))
    return rows
live = loadS("STOCK 1",197,454) + loadS("STOCK 2",600,695)
snap = loadS("Registro de STOCK 1",197,418) + loadS("Registro de STOCK 2",600,695)
def num(x): return float(x) if isinstance(x,(int,float)) else 0.0
L={}; S={}
for r in live: L.setdefault(r['cod'],[]).append(r)
for r in snap: S.setdefault(r['cod'],[]).append(r)
def g(d,k): return sum(num(x['pf']) for x in d.get(k,[]))
def e(d,k): return sum(num(x['pc']) for x in d.get(k,[]))
def f_(d,k): return sum(num(x['costos']) for x in d.get(k,[]))

allk=set(L)|set(S)
nuevos=[]; subieron=[]; bajaron=[]
for k in sorted(allk, key=lambda x:(x is None,x)):
    dg = g(L,k)-g(S,k)
    if k not in S: nuevos.append((k,g(L,k),e(L,k),f_(L,k)))
    elif abs(dg)>0.004: subieron.append((k,g(S,k),g(L,k),dg,e(S,k),e(L,k),f_(S,k),f_(L,k)))
print("### CODIGOS NUEVOS en el rango (no estaban en la foto): %d, sumaG=%.2f sumaE=%.2f sumaF=%.2f"
      % (len(nuevos), sum(x[1] for x in nuevos), sum(x[2] for x in nuevos), sum(x[3] for x in nuevos)))
print("### CODIGOS QUE YA ESTABAN Y CAMBIO SU G: %d, deltaG=%.2f" % (len(subieron), sum(x[3] for x in subieron)))
dE=sum(x[5]-x[4] for x in subieron); dF=sum(x[7]-x[6] for x in subieron)
print("      de los cuales delta E(precio compra)=%.2f  delta F(costos)=%.2f" % (dE,dF))
print()
print("   detalle de los que cambiaron:")
for k,gs,gl,d,es,el,fs,flv in sorted(subieron,key=lambda x:-abs(x[3])):
    print("     %-9s G %8.2f -> %8.2f (%+8.2f)  E %8.2f->%8.2f  F %7.2f->%7.2f" % (k,gs,gl,d,es,el,fs,flv))
print()
tot_new=sum(x[1] for x in nuevos); tot_chg=sum(x[3] for x in subieron)
print("DELTA TOTAL G del rango = %.2f + %.2f = %.2f" % (tot_new,tot_chg,tot_new+tot_chg))
print("Chequeo directo: sumaG live - sumaG snap = %.2f" % (sum(num(r['pf']) for r in live)-sum(num(r['pf']) for r in snap)))

# --- 3. fechas col A de los 20 codigos
lista=['P-00961','P-00962','P-00963','P-00964','P-00967','P-00969','P-00972','P-00975','P-00976','P-00978','P-00979','P-00980','P-00981','P-00982','P-00983','P-00984','P-00986','P-00987','P-00988','P-00996']
print("\n### Fechas col A (STOCK vivo) de los 20 codigos")
mm=collections.Counter()
for k in lista:
    for x in L.get(k,[]):
        f=x['fecha']; key="%04d-%02d"%(f.year,f.month) if isinstance(f,datetime.datetime) else str(f)
        mm[key]+=1
print("   ", dict(mm))

# --- 4. cuanto del delta viene de filas fechadas en agosto
ago_new = [x for x in nuevos]
print("\n### Meses de los codigos NUEVOS")
mn=collections.Counter()
for k,_,_,_ in nuevos:
    for x in L.get(k,[]):
        f=x['fecha']; key="%04d-%02d"%(f.year,f.month) if isinstance(f,datetime.datetime) else str(f)
        mn[key]+=1
print("   ", dict(mn))
print("\n### Meses de los codigos que CAMBIARON")
mc=collections.Counter(); mcv=collections.defaultdict(float)
for row in subieron:
    k=row[0]
    for x in L.get(k,[]):
        f=x['fecha']; key="%04d-%02d"%(f.year,f.month) if isinstance(f,datetime.datetime) else str(f)
        mc[key]+=1; mcv[key]+=row[3]
print("   n:", dict(mc))
print("   deltaG:", {k:round(v,2) for k,v in mcv.items()})
