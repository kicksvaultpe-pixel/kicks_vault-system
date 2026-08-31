# -*- coding: utf-8 -*-
import openpyxl
from datetime import datetime
P = "C:/Users/usuario/Downloads/Registro KV Excel.xlsx"
wbV = openpyxl.load_workbook(P, data_only=True)
wbF = openpyxl.load_workbook(P, data_only=False)
def num(x): return float(x) if isinstance(x,(int,float)) else 0.0

for sh in ["STOCK 1","STOCK 2"]:
    v=wbV[sh]
    filas=[]
    for r in range(2, v.max_row+1):
        cod=v.cell(r,2).value
        if not cod: continue
        f=v.cell(r,1).value
        filas.append((r,cod,f,num(v.cell(r,5).value),num(v.cell(r,6).value),num(v.cell(r,7).value),v.cell(r,8).value))
    print(f"--- {sh}: {len(filas)} filas con codigo; primera r{filas[0][0]} ultima r{filas[-1][0]}")
    ago=[x for x in filas if isinstance(x[2],datetime) and x[2].year==2026 and x[2].month==8]
    print(f"    dateadas AGOSTO-2026: {len(ago)} filas, SUM(G)={sum(x[5] for x in ago):,.2f}  rows {[x[0] for x in ago]}")
    jul=[x for x in filas if isinstance(x[2],datetime) and x[2].year==2026 and x[2].month==7]
    print(f"    dateadas JULIO-2026 : {len(jul)} filas, SUM(G)={sum(x[5] for x in jul):,.2f}")
    print(f"    SUM(G) TODA la hoja = {sum(x[5] for x in filas):,.2f}")
    # rango de la formula
    if sh=="STOCK 1": r0,r1=197,454
    else: r0,r1=600,695
    enr=[x for x in filas if r0<=x[0]<=r1]
    print(f"    en rango {r0}:{r1}: {len(enr)} filas con codigo, SUM(G)={sum(x[5] for x in enr):,.2f}")
    agoen=[x for x in enr if isinstance(x[2],datetime) and x[2].year==2026 and x[2].month==8]
    print(f"       de esas, AGOSTO: {len(agoen)} filas SUM(G)={sum(x[5] for x in agoen):,.2f}")
    fuera=[x for x in ago if not (r0<=x[0]<=r1)]
    print(f"    filas de AGOSTO FUERA del rango de la formula: {len(fuera)} SUM(G)={sum(x[5] for x in fuera):,.2f} rows {[x[0] for x in fuera][:60]}")
