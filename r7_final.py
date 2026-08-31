# -*- coding: utf-8 -*-
import openpyxl, datetime, collections, json, urllib.request
P = "C:/Users/usuario/Downloads/Registro KV Excel.xlsx"
wbv = openpyxl.load_workbook(P, data_only=True)

lista=['P-00961','P-00962','P-00963','P-00964','P-00967','P-00969','P-00972','P-00975','P-00976','P-00978','P-00979','P-00980','P-00981','P-00982','P-00983','P-00984','P-00986','P-00987','P-00988','P-00996']

# (a) 'Registro de STOCK 1' tiene datos mas alla de la fila 418?
ws=wbv["Registro de STOCK 1"]
n=0; tot=0.0
for r in range(419,455):
    v=[ws.cell(r,c).value for c in range(1,9)]
    if any(x is not None for x in v):
        n+=1
        if isinstance(v[6],(int,float)): tot+=float(v[6])
print("### 'Registro de STOCK 1' filas 419-454 (FUERA de su formula SUM(G197:G418)):")
print("    filas con datos = %d, sumaG = %.2f" % (n,tot))
print("    -> la hoja 'foto' SI tiene esas filas; la formula simplemente no las suma.")
ws2=wbv["Registro de STOCK 2"]
n2=0
for r in range(600,696):
    v=[ws2.cell(r,c).value for c in range(1,9)]
    if any(x is not None for x in v): n2+=1
print("    'Registro de STOCK 2' filas 600-695 con datos = %d (live=96)" % n2)
print()

# (b) FLUJO G vs STOCK E por codigo
def loadF(s):
    w=wbv[s]; out={}
    for r in range(2,w.max_row+1):
        cod=w.cell(r,3).value
        acc=w.cell(r,2).value
        if not cod: continue
        if str(acc).strip().upper()!="COMPRA": continue
        out.setdefault(str(cod).strip(),[]).append((r,w.cell(r,7).value,w.cell(r,8).value))
    return out
F=loadF("FLUJO")
def stockE(cod):
    for sh,r0,r1 in (("STOCK 1",197,454),("STOCK 2",600,695)):
        w=wbv[sh]
        for r in range(r0,r1+1):
            if w.cell(r,2).value and str(w.cell(r,2).value).strip()==cod:
                return w.cell(r,5).value
    return None
print("### FLUJO col G (soles) vs STOCK col E (precio compra) y tipo de cambio implicito")
sF=0.0; sE=0.0
for k in lista:
    fg = F[k][0][1]; fu = F[k][0][2]; se = stockE(k)
    tc = (abs(float(fg))/float(fu)) if fg and fu else None
    print("   %-9s FLUJO_G=%-10.2f STOCK_E=%-10.2f iguales=%-5s US$=%-8.2f TC=%.4f"
          % (k,abs(float(fg)),float(se),abs(abs(float(fg))-float(se))<0.005,float(fu),tc))
    sF+=abs(float(fg)); sE+=float(se)
print("   SUMA FLUJO_G=%.2f   SUMA STOCK_E=%.2f   -> es EL MISMO dato en dos hojas" % (sF,sE))
print()

# (c) descomposicion exacta del gap 17049.28
def loadS(sh,r0,r1):
    w=wbv[sh]; out=[]
    for r in range(r0,r1+1):
        v=[w.cell(r,c).value for c in range(1,9)]
        if all(x is None for x in v): continue
        out.append((sh,r,v[0],str(v[1]).strip() if v[1] else None,v[4],v[5],v[6]))
    return out
live = loadS("STOCK 1",197,454)+loadS("STOCK 2",600,695)
snap = loadS("Registro de STOCK 1",197,418)+loadS("Registro de STOCK 2",600,695)
def num(x): return float(x) if isinstance(x,(int,float)) else 0.0
Sm={}
for sh,r,f,c,e,co,pf in snap: Sm[(sh.replace("Registro de ",""),r)]=(e,co,pf)
dE_ago=dE_no=dF_ago=dF_no=0.0
for sh,r,f,c,e,co,pf in live:
    prev=Sm.get((sh,r))
    if prev is None: prev=(0,0,0)
    de=num(e)-num(prev[0]); df=num(co)-num(prev[1])
    esago = isinstance(f,datetime.datetime) and f.year==2026 and f.month==8
    if esago: dE_ago+=de; dF_ago+=df
    else:     dE_no+=de;  dF_no+=df
print("### Descomposicion del delta del rango (live - 'Registro de'), por mes de la col A")
print("   delta E filas de AGOSTO      = %10.2f" % dE_ago)
print("   delta F filas de AGOSTO      = %10.2f" % dF_ago)
print("   delta E filas NO de agosto   = %10.2f   <- el 15,874.08 del hallazgo" % dE_no)
print("   delta F filas NO de agosto   = %10.2f   <- lo que el hallazgo OMITE" % dF_no)
print("   ---------------------------------------------")
print("   GAP TOTAL (no-agosto)        = %10.2f" % (dE_no+dF_no))
print("   comprobacion 21586.66 - 4537.38 = %.2f" % (21586.66-4537.38))
print()

# (d) base de datos
def get(t,q=""):
    out=[]; off=0
    while True:
        req=urllib.request.Request("https://erpkxogskfrtfttawapg.supabase.co/rest/v1/%s?select=*%s"%(t,q))
        req.add_header("apikey","sb_publishable_qgRu9snBNMnDpQ3o9UOR4A_HAzXUhA1")
        req.add_header("Authorization","Bearer sb_publishable_qgRu9snBNMnDpQ3o9UOR4A_HAzXUhA1")
        req.add_header("Range","%d-%d"%(off,off+999))
        d=json.loads(urllib.request.urlopen(req).read().decode())
        out+=d
        if len(d)<1000: break
        off+=1000
    return out
inv=get("inventario")
print("### Base de datos: los 20 codigos en inventario")
im={r['codigo_producto']:r for r in inv}
cnt=collections.Counter()
for k in lista:
    r=im.get(k)
    if not r: print("   %s AUSENTE"%k); continue
    cnt[(r.get('pago_estado'),r.get('tipo_pago'))]+=1
    print("   %-9s fecha=%s pago=%-9s tipo=%-8s S/=%-9s US$=%-8s tc=%-6s final=%s"
          % (k,str(r.get('fecha_compra'))[:10],r.get('pago_estado'),r.get('tipo_pago'),
             r.get('precio_compra_soles'),r.get('precio_compra_usd'),r.get('tipo_cambio'),r.get('precio_final')))
print("   resumen (pago_estado,tipo_pago):",dict(cnt))
borrados=['P-00965','P-00966','P-00968','P-00970','P-00971','P-00973','P-00974','P-00977','P-00985','P-00994','P-00995']
print("\n### Base de datos: los 11 codigos BORRADOS del FLUJO vivo")
for k in borrados:
    r=im.get(k)
    print("   %-9s %s"%(k, "AUSENTE en inventario" if not r else
        "fecha=%s pago=%-9s tipo=%-8s S/=%-9s US$=%s"%(str(r.get('fecha_compra'))[:10],r.get('pago_estado'),r.get('tipo_pago'),r.get('precio_compra_soles'),r.get('precio_compra_usd'))))
