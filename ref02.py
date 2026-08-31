# -*- coding: utf-8 -*-
import openpyxl
from decimal import Decimal, ROUND_HALF_UP

P = "C:/Users/usuario/Downloads/Registro KV Excel.xlsx"
wbv = openpyxl.load_workbook(P, data_only=True)
wbf = openpyxl.load_workbook(P, data_only=False)

def suma(sh, col, r1, r2):
    s = Decimal(0); n = 0; nonnum = []
    ws = wbv[sh]
    for r in range(r1, r2+1):
        v = ws["%s%d" % (col, r)].value
        if isinstance(v, (int, float)):
            s += Decimal(str(v)); n += 1
        elif v is not None:
            nonnum.append((r, v))
    return s, n, nonnum

print("=== SUMAS RECALCULADAS (valores cacheados del archivo) ===")

s1, n1, nn1 = suma("STOCK 1", "G", 197, 454)
print("SUM('STOCK 1'!G197:G454)          = %s   (%d celdas numericas)" % (s1, n1))
if nn1: print("   no numericas:", nn1[:10])

s2, n2, nn2 = suma("STOCK 2", "G", 600, 695)
print("SUM('STOCK 2'!G600:G695)          = %s   (%d celdas numericas)" % (s2, n2))
if nn2: print("   no numericas:", nn2[:10])

s3, n3, nn3 = suma("Registro de STOCK 1", "G", 197, 418)
print("SUM('Registro de STOCK 1'!G197:G418) = %s   (%d celdas)" % (s3, n3))
if nn3: print("   no numericas:", nn3[:10])

s4, n4, nn4 = suma("Registro de STOCK 2", "G", 600, 695)
print("SUM('Registro de STOCK 2'!G600:G695) = %s   (%d celdas)" % (s4, n4))
if nn4: print("   no numericas:", nn4[:10])

print()
print("=== CHEQUEO DE FORMULAS ===")
print("STOCK1 O455: %s - 84285.52 = %s   (cacheado 17319.09)  match=%s" %
      (s1, s1-Decimal("84285.52"), (s1-Decimal("84285.52")).quantize(Decimal("0.01"))==Decimal("17319.09")))
print("STOCK1 O456: %s - 234 = %s   (cacheado 4267.57)  match=%s" %
      (s2, s2-Decimal("234"), (s2-Decimal("234")).quantize(Decimal("0.01"))==Decimal("4267.57")))
print("Reg O419: %s - 72000 - 135.87 = %s   (cacheado 15584.32)  match=%s" %
      (s3, s3-Decimal("72135.87"), (s3-Decimal("72135.87")).quantize(Decimal("0.01"))==Decimal("15584.32")))
print("Reg O420 formula apunta a %r" % wbf["Registro de STOCK 1"]["O420"].value)
print("   SUM('STOCK 2'!G600:G695) actual = %s ; cacheado O420 = 4501.57" % s2)
print("   SUM('Registro de STOCK 2'!G600:G695) = %s" % s4)

print()
print("=== DIFERENCIAS ===")
print("87720.19 vs 84285.52 -> %s" % (s3 - Decimal("84285.52")))
print("s3 real = %s" % s3)
