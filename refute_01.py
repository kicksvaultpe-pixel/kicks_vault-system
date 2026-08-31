import openpyxl, datetime
from decimal import Decimal

P = "C:/Users/usuario/Downloads/Registro KV Excel.xlsx"
wbv = openpyxl.load_workbook(P, data_only=True)
wbf = openpyxl.load_workbook(P, data_only=False)

def dump(ws_name, r1, r2, wb):
    ws = wb[ws_name]
    rows = []
    for r in range(r1, r2+1):
        a = ws.cell(r, 1).value
        b = ws.cell(r, 2).value
        c = ws.cell(r, 3).value
        d = ws.cell(r, 4).value
        e = ws.cell(r, 5).value
        f = ws.cell(r, 6).value
        g = ws.cell(r, 7).value
        h = ws.cell(r, 8).value
        rows.append((r,a,b,c,d,e,f,g,h))
    return rows

# 1. The formula cells themselves
print("=== STOCK 1 N/O 450-460 formulas ===")
ws = wbf["STOCK 1"]
wsv = wbv["STOCK 1"]
for r in range(448, 462):
    for col in ("M","N","O","P"):
        fv = ws[f"{col}{r}"].value
        vv = wsv[f"{col}{r}"].value
        if fv is not None or vv is not None:
            print(f"  {col}{r}: formula={fv!r}  value={vv!r}")

print()
print("=== Registro de STOCK 1 N/O 415-425 ===")
ws2 = wbf["Registro de STOCK 1"]
wsv2 = wbv["Registro de STOCK 1"]
for r in range(413, 426):
    for col in ("M","N","O","P"):
        fv = ws2[f"{col}{r}"].value
        vv = wsv2[f"{col}{r}"].value
        if fv is not None or vv is not None:
            print(f"  {col}{r}: formula={fv!r}  value={vv!r}")

print()
print("=== dimensions ===")
for n in ["STOCK 1","STOCK 2","Registro de STOCK 1","Registro de STOCK 2"]:
    w = wbv[n]
    print(f"  {n}: max_row={w.max_row} max_col={w.max_column}")
