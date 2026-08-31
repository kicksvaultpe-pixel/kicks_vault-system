# -*- coding: utf-8 -*-
import openpyxl, datetime, json, urllib.request

P = "C:/Users/usuario/Downloads/Registro KV Excel.xlsx"
wbv = openpyxl.load_workbook(P, data_only=True)
vs = wbv['Registro de STOCK 2']; vl = wbv['STOCK 2']

print("=== DESCOMPOSICION DEL 614.13 (snapshot G600:G695) ===")
pre_ago = []; ago = []
for r in range(600, 696):
    g = vs.cell(r,7).value
    if not isinstance(g,(int,float)) or g == 0: continue
    a = vs.cell(r,1).value
    b = vs.cell(r,2).value
    mes = a.strftime("%Y-%m") if isinstance(a, datetime.datetime) else "?"
    (ago if mes=="2026-08" else pre_ago).append((r,b,mes,g))

s_pre = sum(x[3] for x in pre_ago); s_ago = sum(x[3] for x in ago)
print(" NO-AGOSTO:")
for x in pre_ago: print(f"   r{x[0]} {x[1]} {x[2]} G={x[3]!r}")
print(f"   SUBTOTAL NO-AGOSTO = {s_pre!r}  -> {s_pre:.2f}")
print(" AGOSTO:")
for x in ago: print(f"   r{x[0]} {x[1]} {x[2]} G={x[3]!r}")
print(f"   SUBTOTAL AGOSTO = {s_ago!r} -> {s_ago:.2f}")
print(f"\n  subtotal_no_agosto == 234 ?  {s_pre == 234.0}   (dif {s_pre-234.0!r})")
print(f"  614.13 = {s_pre:.2f} + {s_ago:.2f} = {s_pre+s_ago:.2f}")

print("\n=== Valores LIVE de esos mismos productos (pre-agosto) ===")
tot_live_pre = 0.0
for r,b,mes,g in pre_ago:
    lg = vl.cell(r,7).value
    tot_live_pre += lg
    print(f"   r{r} {b} {mes}  SNAP={g:>8.2f}  LIVE={lg:>8.2f}")
print("  ...mas filas pre-agosto con SNAP=0:")
for r in range(600, 609):
    g = vs.cell(r,7).value
    if isinstance(g,(int,float)) and g == 0:
        lg = vl.cell(r,7).value
        tot_live_pre += lg
        a = vl.cell(r,1).value
        print(f"   r{r} {vl.cell(r,2).value} {a.strftime('%Y-%m-%d') if isinstance(a,datetime.datetime) else a}  SNAP=0.00  LIVE={lg:>8.2f}")
print(f"  TOTAL LIVE de filas pre-agosto = {tot_live_pre:.2f}")
print(f"  4501.57 - {tot_live_pre:.2f} = {4501.57-tot_live_pre:.2f}")

print("\n=== ULTIMA FILA CON CODIGO EN CADA HOJA (rango 600-695) ===")
for nm, ws in (("Registro de STOCK 2", vs), ("STOCK 2", vl)):
    last = None
    for r in range(600, 696):
        if ws.cell(r,2).value: last = r
    a = ws.cell(last,1).value
    if isinstance(a,(int,float)):
        a = (datetime.datetime(1899,12,30)+datetime.timedelta(days=a)).strftime("%Y-%m-%d")+" (serial)"
    elif isinstance(a, datetime.datetime): a = a.strftime("%Y-%m-%d")
    print(f"  {nm}: ultima fila con codigo = {last}, cod={ws.cell(last,2).value}, fecha={a}")

print("\n=== BUSCAR P-00961 y P-01000 en hoja COMPRA (col A) ===")
co = wbv['COMPRA']
want = {"P-00961","P-00962","P-00963","P-00975","P-00987","P-00994","P-01000"}
found = {}
for row in co.iter_rows(min_row=1, max_row=co.max_row, max_col=6, values_only=True):
    if row and row[0] in want:
        found[row[0]] = row
for k in sorted(want):
    print(f"  {k}: {'NO ESTA EN COMPRA' if k not in found else found[k]}")

print("\n=== BASE DE DATOS: inventario P-01000 y P-00961 ===")
H = {"apikey":"sb_publishable_qgRu9snBNMnDpQ3o9UOR4A_HAzXUhA1",
     "Authorization":"Bearer sb_publishable_qgRu9snBNMnDpQ3o9UOR4A_HAzXUhA1"}
for cod in ("P-01000","P-00961","P-00975","P-00987"):
    u = f"https://erpkxogskfrtfttawapg.supabase.co/rest/v1/inventario?codigo_producto=eq.{cod}&select=codigo_producto,fecha_compra,modelo,estado,pago_estado,tipo_pago,precio_compra_soles,precio_final"
    req = urllib.request.Request(u, headers=H)
    try:
        print("  ", json.loads(urllib.request.urlopen(req, timeout=30).read().decode()))
    except Exception as e:
        print("   ERR", cod, e)
