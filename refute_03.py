import openpyxl, datetime
from collections import defaultdict

P = "C:/Users/usuario/Downloads/Registro KV Excel.xlsx"
wbv = openpyxl.load_workbook(P, data_only=True)

def mk(v):
    if isinstance(v, datetime.datetime): return (v.year, v.month)
    if isinstance(v, datetime.date):     return (v.year, v.month)
    return None

def rows(ws, r1, r2):
    w = wbv[ws]; out={}
    for r in range(r1, r2+1):
        cod = w.cell(r,2).value
        g   = w.cell(r,7).value
        if not isinstance(g,(int,float)): continue
        out[r] = dict(row=r, fecha=w.cell(r,1).value, cod=cod,
                      desc=w.cell(r,3).value, talla=w.cell(r,4).value,
                      e=w.cell(r,5).value, f=w.cell(r,6).value, g=g,
                      estado=w.cell(r,8).value)
    return out

S1  = rows("STOCK 1",197,454)
R1  = rows("Registro de STOCK 1",197,418)
S2  = rows("STOCK 2",600,695)
R2  = rows("Registro de STOCK 2",600,695)

# ---------- A) date grouping of the CURRENT ranges ----------
print("="*70)
print("A) CURRENT ranges G, grouped by month of column A")
for nm, D in (("STOCK 1 G197:G454",S1),("STOCK 2 G600:G695",S2)):
    agg=defaultdict(lambda:[0.0,0])
    for d in D.values():
        k=mk(d["fecha"]); agg[k][0]+=d["g"]; agg[k][1]+=1
    print(f"  -- {nm}")
    for k in sorted(agg, key=lambda x:(x is None, x)):
        print(f"       {str(k):<12} {agg[k][0]:>12,.2f}  n={agg[k][1]}")
aug1 = round(sum(d["g"] for d in S1.values() if mk(d["fecha"])==(2026,8)),2)
aug2 = round(sum(d["g"] for d in S2.values() if mk(d["fecha"])==(2026,8)),2)
n1 = sum(1 for d in S1.values() if mk(d["fecha"])==(2026,8))
n2 = sum(1 for d in S2.values() if mk(d["fecha"])==(2026,8))
print(f"\n  AUG-dated total = {aug1:,.2f} ({n1} rows S1) + {aug2:,.2f} ({n2} rows S2) = {round(aug1+aug2,2):,.2f}   [app says 4,537.38]")

# ---------- B) what actually ENTERED the sheets during august ----------
print()
print("="*70)
print("B) Rows present NOW but NOT in the july snapshot (true 'entered during august')")

def keyof(d):
    return (str(d["cod"]).strip() if d["cod"] is not None else None,
            str(d["desc"]).strip() if d["desc"] is not None else None,
            str(d["talla"]).strip() if d["talla"] is not None else None)

for nm, CUR, PREV in (("STOCK 1",S1,R1), ("STOCK 2",S2,R2)):
    prev_keys=defaultdict(int)
    for d in PREV.values(): prev_keys[keyof(d)] += 1
    new=[]; matched=[]
    for d in sorted(CUR.values(), key=lambda x:x["row"]):
        k=keyof(d)
        if prev_keys.get(k,0)>0:
            prev_keys[k]-=1; matched.append(d)
        else:
            new.append(d)
    tot_new=round(sum(d["g"] for d in new),2)
    print(f"  -- {nm}: {len(new)} NEW rows, G = {tot_new:,.2f}   ({len(matched)} carried over)")
    agg=defaultdict(lambda:[0.0,0])
    for d in new:
        k=mk(d["fecha"]); agg[k][0]+=d["g"]; agg[k][1]+=1
    for k in sorted(agg, key=lambda x:(x is None, x)):
        print(f"       new rows dated {str(k):<12} {agg[k][0]:>11,.2f}  n={agg[k][1]}")
    globals()["NEW_"+nm.replace(" ","")]=new

allnew = NEW_STOCK1 + NEW_STOCK2
print(f"\n  TOTAL new-row G across both = {round(sum(d['g'] for d in allnew),2):,.2f}")
pre_aug = round(sum(d["g"] for d in allnew if mk(d["fecha"]) not in [(2026,8)]),2)
in_aug  = round(sum(d["g"] for d in allnew if mk(d["fecha"])==(2026,8)),2)
print(f"    of which dated BEFORE august (or blank) = {pre_aug:,.2f}   [analyst claims 17,049.28]")
print(f"    of which dated IN august                = {in_aug:,.2f}")

# ---------- C) value changes on carried-over rows ----------
print()
print("="*70)
print("C) Did G change on rows that already existed in july? (costs added later)")
for nm, CUR, PREV in (("STOCK 1",S1,R1), ("STOCK 2",S2,R2)):
    pm=defaultdict(list)
    for d in PREV.values(): pm[keyof(d)].append(d)
    delta=0.0; ch=[]
    for d in CUR.values():
        k=keyof(d)
        if pm.get(k):
            p=pm[k].pop(0)
            if abs((p["g"] or 0)-(d["g"] or 0))>0.004:
                delta += d["g"]-p["g"]; ch.append((d,p))
    print(f"  -- {nm}: {len(ch)} rows changed G, net delta = {round(delta,2):,.2f}")
    for d,p in ch[:15]:
        print(f"       r{d['row']:<5} {str(d['cod'])[:10]:<10} {str(d['desc'])[:26]:<26} G {p['g']:>9,.2f} -> {d['g']:>9,.2f}  ({d['g']-p['g']:+,.2f})  fecha={d['fecha']}")
