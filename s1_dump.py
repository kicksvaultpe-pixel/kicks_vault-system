import openpyxl
p="C:/Users/usuario/Downloads/Registro KV Excel.xlsx"
wbv=openpyxl.load_workbook(p,data_only=True)
wbf=openpyxl.load_workbook(p,data_only=False)
print("SHEETS:", wbv.sheetnames)
for sh in ['STOCK 1','STOCK 2','Registro de STOCK 1','Registro de STOCK 2']:
    w=wbv[sh]
    print(sh,"dims",w.dimensions,"max_row",w.max_row)
