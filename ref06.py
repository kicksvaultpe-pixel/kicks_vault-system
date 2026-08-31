# -*- coding: utf-8 -*-
import openpyxl
from decimal import Decimal
P="C:/Users/usuario/Downloads/Registro KV Excel.xlsx"
wbv=openpyxl.load_workbook(P,data_only=True)
S=wbv["STOCK 1"]; R=wbv["Registro de STOCK 1"]
S2=wbv["STOCK 2"]; R2=wbv["Registro de STOCK 2"]

# 1) columnas A y B identicas entre foto y hoja viva?
bad=[]
for r in range(197,419):
    if S.cell(row=r,column=1).value != R.cell(row=r,column=1).value: bad.append(("A",r))
    if S.cell(row=r,column=2).value != R.cell(row=r,column=2).value: bad.append(("B",r))
print("STOCK1 vs REG: celdas A/B distintas en 197:418 ->", bad if bad else "NINGUNA (mismo set de filas, misma fecha, mismo codigo)")

bad2=[]
for r in range(600,696):
    if S2.cell(row=r,column=1).value != R2.cell(row=r,column=1).value: bad2.append(("A",r))
print("STOCK2 vs REG2: celdas A distintas ->", len(bad2), bad2[:12])

# 2) que columnas cambian
from collections import Counter
ch=Counter()
for r in range(197,419):
    for c,nm in [(5,"E precio compra"),(6,"F costos"),(7,"G final"),(8,"H estado")]:
        if S.cell(row=r,column=c).value != R.cell(row=r,column=c).value: ch[nm]+=1
print("columnas que cambian entre foto y viva (197:418):", dict(ch))

# 3) reconstruir las dos constantes desde cero
def const(Rws, r1, r2):
    tot=Decimal(0); aug=Decimal(0)
    for r in range(r1,r2+1):
        f=Rws.cell(row=r,column=1).value
        v=Rws.cell(row=r,column=7).value
        g=Decimal(str(v)) if isinstance(v,(int,float)) else Decimal(0)
        tot+=g
        if hasattr(f,'year') and (f.year,f.month)==(2026,8): aug+=g
    return tot,aug,tot-aug
print()
for nm,ws,r1,r2,esperado in [("Registro de STOCK 1",R,197,418,"84285.52"),
                              ("Registro de STOCK 2",R2,600,695,"234")]:
    t,a,c = const(ws,r1,r2)
    print("%-22s total=%9s  agosto=%8s  total-agosto=%9s  constante en formula=%s  MATCH=%s"
          % (nm,t,a,c,esperado,c==Decimal(esperado)))
