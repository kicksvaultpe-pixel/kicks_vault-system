# -*- coding: utf-8 -*-
import openpyxl
from datetime import datetime
P = "C:/Users/usuario/Downloads/Registro KV Excel.xlsx"
wbV = openpyxl.load_workbook(P, data_only=True)
def num(x): return float(x) if isinstance(x,(int,float)) else None

# indice COMPRA por codigo
cv=wbV["COMPRA"]; COMPRA={}
for r in range(2,cv.max_row+1):
    cod=cv.cell(r,1).value
    if cod: COMPRA[cod]=dict(row=r,fecha=cv.cell(r,2).value,soles=num(cv.cell(r,6).value),usd=num(cv.cell(r,7).value))

ovr = {  # codigo: monto override en E
 "P-00964":1334.26,"P-00967":1431.30,"P-00969":1431.30,"P-00972":1431.36,"P-00978":518.50,
 "P-00979":526.49,"P-00980":452.72,"P-00981":1351.88,"P-00982":231.27,"P-00983":248.30,
 "P-00986":731.04,"P-00988":1516.02,"P-00996":841.50,
 "P-00961":454.99,"P-00962":530.26,"P-00963":460.47,"P-00975":396.49,"P-00976":399.91,
 "P-00984":647.93,"P-00987":938.09}
print(f"{'COD':10}{'fCompra':12}{'soles':>10}{'usd':>10}{'override':>10}{'ratio':>8}")
tot=0; sinsoles=0
for c,m in ovr.items():
    x=COMPRA.get(c)
    if not x: print(c,"NO ESTA EN COMPRA"); continue
    s=x['soles']; u=x['usd']
    ratio = m/u if u else None
    if s is None: sinsoles+=1
    print(f"{c:10}{str(x['fecha'])[:10]:12}{('VACIO' if s is None else f'{s:.2f}'):>10}{(f'{u:.2f}' if u else '-'):>10}{m:>10.2f}{(f'{ratio:.4f}' if ratio else '-'):>8}")
    tot+=m
print(f"TOTAL OVERRIDES = {tot:,.2f}   con Monto(S/.) VACIO en COMPRA: {sinsoles}/{len(ovr)}")

# CONTRA-PRUEBA: en el rango, cuantas filas tienen E literal y COMPRA CON soles?
print("\n=== Contraprueba: filas del rango con E literal ===")
import openpyxl as ox
wbF = ox.load_workbook(P, data_only=False)
for sh,r0,r1 in [("STOCK 1",197,454),("STOCK 2",600,695)]:
    f=wbF[sh]; v=wbV[sh]
    lit_consoles=0; lit_sinsoles=0; form=0; lit_nocompra=0
    for r in range(r0,r1+1):
        cod=v.cell(r,2).value
        if not cod: continue
        ef=f.cell(r,5).value
        if isinstance(ef,str) and ef.startswith("="): form+=1; continue
        x=COMPRA.get(cod)
        if not x: lit_nocompra+=1
        elif x['soles'] is None: lit_sinsoles+=1
        else: lit_consoles+=1
    print(f"  {sh}: E formula={form}  E literal con soles en COMPRA={lit_consoles}  "
          f"E literal SIN soles (override)={lit_sinsoles}  E literal sin fila COMPRA={lit_nocompra}")

# filas del rango con E formula=0 y COMPRA en USD (aun NO pagadas)
print("\n=== Filas del rango con E=formula(0) y compra en USD (pendientes) ===")
for sh,r0,r1 in [("STOCK 1",197,454),("STOCK 2",600,695)]:
    f=wbF[sh]; v=wbV[sh]; n=0; usd=0
    for r in range(r0,r1+1):
        cod=v.cell(r,2).value
        if not cod: continue
        ef=f.cell(r,5).value
        if not (isinstance(ef,str) and ef.startswith("=")): continue
        x=COMPRA.get(cod)
        if x and x['soles'] is None and x['usd']:
            n+=1; usd+=x['usd']
    print(f"  {sh}: {n} filas, US$ {usd:,.2f} sin convertir a soles todavia")
