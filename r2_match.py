# -*- coding: utf-8 -*-
import openpyxl, datetime, collections
P = "C:/Users/usuario/Downloads/Registro KV Excel.xlsx"
wb = openpyxl.load_workbook(P, data_only=True)

def load(sheetname):
    ws = wb[sheetname]
    rows=[]
    for r in range(2, ws.max_row+1):
        v=[ws.cell(r,c).value for c in range(1,10)]
        if all(x is None for x in v[:8]): continue
        rows.append(dict(row=r,fecha=v[0],accion=v[1],cod=v[2],desc=v[3],talla=v[4],
                         ing=v[5],egr=v[6],usd=v[7],info=v[8]))
    return rows

fl=load("FLUJO"); rg=load("Registro de FLUJO")
def isC(r): return str(r['accion']).strip().upper()=="COMPRA" if r['accion'] else False
def m(v): return "%04d-%02d"%(v.year,v.month) if isinstance(v,datetime.datetime) else None

flC=[r for r in fl if isC(r)]
rgC=[r for r in rg if isC(r)]
fl_jj=[r for r in flC if m(r['fecha']) in ("2026-06","2026-07")]
rg_jj=[r for r in rgC if m(r['fecha']) in ("2026-06","2026-07")]
print("FLUJO vivo COMPRA jun+jul:", len(fl_jj), "sum soles=%.2f"%sum(float(r['egr']) for r in fl_jj if r['egr'] is not None))
print("REG      COMPRA jun+jul:", len(rg_jj), "sum soles=%.2f"%sum(float(r['egr']) for r in rg_jj if r['egr'] is not None))

# indices por codigo
def idx(rows):
    d=collections.defaultdict(list)
    for r in rows: d[str(r['cod']).strip() if r['cod'] else None].append(r)
    return d
fi=idx(fl_jj); ri=idx(rg_jj)
print("codigos distintos FLUJO jun-jul:", len(fi), " REG jun-jul:", len(ri))
print("codigos en FLUJO no en REG:", sorted(k for k in fi if k not in ri))
print("codigos en REG no en FLUJO:", sorted(k for k in ri if k not in fi))

print()
print("%-10s %-22s %-12s %-12s %-10s %-10s" % ("COD","FECHA_FLUJO","S/_FLUJO","S/_REG","US$_FLUJO","US$_REG"))
tot_new=0.0; tot_chg=0.0; nuevos=[]; cambios=[]; iguales=[]
for k in sorted(fi):
    a=fi[k][0]
    b=ri[k][0] if k in ri else None
    sa=a['egr']; sb=b['egr'] if b else "SIN-FILA"
    ua=a['usd']; ub=b['usd'] if b else "-"
    print("%-10s %-22s %-12s %-12s %-10s %-10s" % (k, str(a['fecha'])[:19], sa, sb, ua, ub))
    if b is None:
        if sa is not None: tot_new+=float(sa); nuevos.append((k,float(sa)))
    elif sb is None and sa is not None:
        tot_new+=float(sa); nuevos.append((k,float(sa)))
    elif sa is not None and sb is not None and abs(float(sa)-float(sb))>0.004:
        tot_chg+=float(sa)-float(sb); cambios.append((k,float(sb),float(sa)))
    else:
        iguales.append((k,sa))
print()
print("SOLES QUE APARECIERON (None/sin-fila -> valor): %.2f  en %d codigos" % (tot_new,len(nuevos)))
for k,v in nuevos: print("   nuevo %s %.2f"%(k,v))
print("SOLES QUE CAMBIARON DE VALOR: delta %.2f en %d codigos" % (tot_chg,len(cambios)))
for k,b,a in cambios: print("   chg %s  %.2f -> %.2f"%(k,b,a))
print("SOLES YA PRESENTES E IGUALES: %d codigos, suma %.2f" % (len(iguales), sum(float(v) for k,v in iguales if v is not None)))
for k,v in iguales: print("   igual %s %s"%(k,v))
