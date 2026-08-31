# -*- coding: utf-8 -*-
import openpyxl
from decimal import Decimal
P = "C:/Users/usuario/Downloads/Registro KV Excel.xlsx"
wbv = openpyxl.load_workbook(P, data_only=True)
S=wbv["STOCK 1"]; R=wbv["Registro de STOCK 1"]
S2=wbv["STOCK 2"]; R2=wbv["Registro de STOCK 2"]
def d(v): return Decimal(str(v)) if isinstance(v,(int,float)) else Decimal(0)

print("=== LAS 16 FILAS FECHADAS AGOSTO DENTRO DE 'Registro de STOCK 1' G197:G418 ===")
tot=Decimal(0); totS=Decimal(0)
for r in range(197,419):
    f=R.cell(row=r,column=1).value
    if hasattr(f,'year') and (f.year,f.month)==(2026,8):
        gr=d(R.cell(row=r,column=7).value); gs=d(S.cell(row=r,column=7).value)
        tot+=gr; totS+=gs
        print("  f%-4d %s %-10s G_reg=%9s  G_stock1=%9s  delta=%8s" %
              (r, f.date(), R.cell(row=r,column=2).value, gr, gs, gs-gr))
print("  ---- SUM G_reg (agosto) = %s   SUM G_stock1 (agosto) = %s" % (tot, totS))

print()
print("=== PRUEBA CENTRAL ===")
base_snapshot = Decimal("87720.19")
print("  SUM('Registro de STOCK 1'!G197:G418)                 = %s" % base_snapshot)
print("  menos filas fechadas AGOSTO dentro de esa foto       = %s" % tot)
print("  = %s   <-- LA CONSTANTE DE LA FORMULA ES 84285.52" % (base_snapshot-tot))
print("  coincide exacto: %s" % ((base_snapshot-tot)==Decimal("84285.52")))

print()
print("=== DESCOMPOSICION DE 17319.09 ===")
print("  revaluacion de filas ya existentes (101604.61-87720.19) = 13884.42")
print("  + valor de las 16 filas de agosto ya cargado en la foto = %s" % tot)
print("  = %s   (cacheado O455 = 17319.09)" % (Decimal("13884.42")+tot))

print()
print("=== MISMO PATRON EN STOCK 2 (constante 234) ===")
regtot=Decimal(0); regaug=Decimal(0); n=0
for r in range(600,696):
    f=R2.cell(row=r,column=1).value
    g=d(R2.cell(row=r,column=7).value)
    regtot+=g
    if hasattr(f,'year') and (f.year,f.month)==(2026,8): regaug+=g; n+=1
print("  SUM('Registro de STOCK 2'!G600:G695)      = %s" % regtot)
print("  de eso, filas fechadas AGOSTO             = %s (%d filas)" % (regaug,n))
print("  regtot - regaug                           = %s   (constante en formula = 234)" % (regtot-regaug))

print()
print("=== FILAS AGOSTO EN LAS HOJAS VIVAS (para el 4537.38 de la app) ===")
a1=Decimal(0); c1=0
for r in range(197,455):
    f=S.cell(row=r,column=1).value
    if hasattr(f,'year') and (f.year,f.month)==(2026,8): a1+=d(S.cell(row=r,column=7).value); c1+=1
a2=Decimal(0); c2=0
for r in range(600,696):
    f=S2.cell(row=r,column=1).value
    if hasattr(f,'year') and (f.year,f.month)==(2026,8): a2+=d(S2.cell(row=r,column=7).value); c2+=1
print("  STOCK 1 agosto = %s (%d filas)" % (a1,c1))
print("  STOCK 2 agosto = %s (%d filas)" % (a2,c2))
print("  total          = %s (%d filas)" % (a1+a2, c1+c2))
