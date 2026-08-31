# -*- coding: utf-8 -*-
import openpyxl, datetime
p = "C:/Users/usuario/Downloads/Registro KV Excel.xlsx"
wv = openpyxl.load_workbook(p, data_only=True)
def num(x): return float(x) if isinstance(x,(int,float)) else 0.0

CODES = ["P-00997","P-00998","P-00999","P-01000","P-01013"]
SHEETS = ["STOCK 1","STOCK 2","Registro de STOCK 1","Registro de STOCK 2"]

print("### Ubicacion de los 5 codigos en cada hoja ###")
found = {}
for sh in SHEETS:
    ws = wv[sh]
    for r in range(1, ws.max_row+1):
        b = ws[f"B{r}"].value
        if isinstance(b,str) and b.strip() in CODES:
            a = ws[f"A{r}"].value
            fecha = a.date() if isinstance(a,datetime.datetime) else a
            found.setdefault(b.strip(),[]).append(
                (sh, r, fecha, num(ws[f"E{r}"].value), num(ws[f"F{r}"].value),
                 num(ws[f"G{r}"].value), ws[f"H{r}"].value))
for c in CODES:
    print(f"\n{c}:")
    for sh, r, fecha, e, f_, g, h in found.get(c, []):
        inrange = ("EN RANGO" if (sh in ("STOCK 1","Registro de STOCK 1") and 197<=r<=(454 if sh=="STOCK 1" else 418))
                   or (sh in ("STOCK 2","Registro de STOCK 2") and 600<=r<=695) else "fuera")
        print(f"   {sh:22s} fila {r:4d} fecha={fecha} E={e:9,.2f} F={f_:8,.2f} G={g:9,.2f} H={h!r}  [{inrange}]")
    if not found.get(c): print("   NO ENCONTRADO")

print()
print("### Suma de G de esos codigos ###")
for grp, codes in [("STOCK1-4", ["P-00997","P-00998","P-00999","P-01013"]), ("STOCK2-1",["P-01000"])]:
    for sh in SHEETS:
        s = sum(g for c in codes for (s2,r,fe,e,f_,g,h) in found.get(c,[]) if s2==sh)
        if s: print(f"  {grp} en {sh:22s} = {s:,.2f}")
