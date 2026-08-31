# -*- coding: utf-8 -*-
import openpyxl, datetime
P = "C:/Users/usuario/Downloads/Registro KV Excel.xlsx"
wbv = openpyxl.load_workbook(P, data_only=True)
wbf = openpyxl.load_workbook(P, data_only=False)

def d(x):
    return x.strftime("%Y-%m-%d") if isinstance(x, datetime.datetime) else (str(x) if x is not None else "")

print("### FORMULAS crudas en 'Registro de STOCK 2' filas 600-610 (A..H)")
fs = wbf['Registro de STOCK 2']
for r in range(600, 611):
    print("  r%d: %s" % (r, [fs.cell(row=r, column=c).value for c in range(1, 9)]))

print("\n### FORMULAS crudas en 'STOCK 2' filas 600-610 (A..H)")
fl = wbf['STOCK 2']
for r in range(600, 611):
    print("  r%d: %s" % (r, [fl.cell(row=r, column=c).value for c in range(1, 9)]))

print("\n\n### COMPARACION FILA A FILA 600-695  (SNAP=Registro de STOCK 2, LIVE=STOCK 2)")
vs = wbv['Registro de STOCK 2']; vl = wbv['STOCK 2']
print(f"{'fil':>4} | {'SNAP fecha':10} {'cod':>9} {'G':>9} {'estado':11} | {'LIVE fecha':10} {'cod':>9} {'G':>9} {'estado':11} | igual?")
ts = tl = 0.0
for r in range(600, 696):
    sa, sb, sg, sh = vs.cell(r,1).value, vs.cell(r,2).value, vs.cell(r,7).value, vs.cell(r,8).value
    la, lb, lg, lh = vl.cell(r,1).value, vl.cell(r,2).value, vl.cell(r,7).value, vl.cell(r,8).value
    sgn = sg if isinstance(sg,(int,float)) else 0
    lgn = lg if isinstance(lg,(int,float)) else 0
    ts += sgn; tl += lgn
    same = "SI" if (str(sb)==str(lb) and abs(sgn-lgn)<0.005) else "<<< DIF"
    print(f"{r:>4} | {d(sa):10} {str(sb):>9} {sgn:>9.2f} {str(sh)[:11]:11} | {d(la):10} {str(lb):>9} {lgn:>9.2f} {str(lh)[:11]:11} | {same}")
print(f"\n TOTAL SNAP={ts:.2f}   TOTAL LIVE={tl:.2f}   dif={tl-ts:.2f}")
