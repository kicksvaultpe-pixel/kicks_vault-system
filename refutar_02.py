# -*- coding: utf-8 -*-
import openpyxl

P = "C:/Users/usuario/Downloads/Registro KV Excel.xlsx"
wb = openpyxl.load_workbook(P, data_only=True)

def s(sheet, col, r1, r2):
    sh = wb[sheet]
    tot = 0.0
    n = 0
    for r in range(r1, r2 + 1):
        v = sh[f"{col}{r}"].value
        if isinstance(v, (int, float)):
            tot += float(v)
            n += 1
    return round(tot, 2), n

print("=== SUMAS CRUDAS QUE YO MISMO CALCULO ===")
a = s("STOCK 1", "G", 197, 454)
b = s("Registro de STOCK 1", "G", 197, 418)
c = s("STOCK 2", "G", 600, 695)
d = s("Registro de STOCK 2", "G", 600, 695)
print(f"SUM('STOCK 1'!G197:G454)             = {a[0]:>12,.2f}  ({a[1]} celdas numericas)")
print(f"SUM('Registro de STOCK 1'!G197:G418) = {b[0]:>12,.2f}  ({b[1]} celdas numericas)")
print(f"SUM('STOCK 2'!G600:G695)             = {c[0]:>12,.2f}  ({c[1]} celdas numericas)")
print(f"SUM('Registro de STOCK 2'!G600:G695) = {d[0]:>12,.2f}  ({d[1]} celdas numericas)")
print()
print(f"Analista dice: 101604.61 / 87720.19 / 4501.57 / 614.13")
print(f"delta1 = {a[0]-b[0]:,.2f}   delta2 = {c[0]-d[0]:,.2f}   total = {(a[0]-b[0])+(c[0]-d[0]):,.2f}")
print()

# Cuantas filas de STOCK 2 tiene el Registro de STOCK 2 realmente?
r2 = wb["Registro de STOCK 2"]
s2 = wb["STOCK 2"]
print("=== Ultima fila con datos ===")
for name in ["STOCK 1", "Registro de STOCK 1", "STOCK 2", "Registro de STOCK 2"]:
    sh = wb[name]
    last = 0
    for r in range(1, sh.max_row + 1):
        if sh[f"B{r}"].value not in (None, ""):
            last = r
    print(f"  {name:<24} max_row={sh.max_row:<6} ultima fila con CODIGO en B = {last}")
print()

# ALINEACION: comparar codigo por fila en el rango 600-695
print("=== ALINEACION DE FILAS 600-695: STOCK 2 vs Registro de STOCK 2 ===")
desal = 0
for r in range(600, 696):
    c1 = s2[f"B{r}"].value
    c2 = r2[f"B{r}"].value
    g1 = s2[f"G{r}"].value
    g2 = r2[f"G{r}"].value
    if c1 != c2:
        desal += 1
    if r <= 615 or c1 != c2:
        print(f"  fila {r}: STOCK2 B={str(c1):<12} G={g1!r:<12} | REG B={str(c2):<12} G={g2!r}")
print(f"  filas desalineadas por codigo en 600-695: {desal}")
