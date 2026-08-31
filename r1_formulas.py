# -*- coding: utf-8 -*-
import openpyxl

P = "C:/Users/usuario/Downloads/Registro KV Excel.xlsx"

wf = openpyxl.load_workbook(P, data_only=False)
wv = openpyxl.load_workbook(P, data_only=True)

print("=== HOJAS ===")
print(wf.sheetnames)

def dump(sheet, cells, label):
    print("\n=== %s (%s) ===" % (label, sheet))
    sf = wf[sheet]; sv = wv[sheet]
    for c in cells:
        print("  %-6s formula=%-60s valor=%s" % (c, repr(sf[c].value), repr(sv[c].value)))

# Bloque STOCK 1 N455:O457
dump("STOCK 1", ["N455","O455","N456","O456","N457","O457"], "COMPRAS MES actual")
# Bloque Registro de STOCK 1 N419:O421
dump("Registro de STOCK 1", ["N419","O419","N420","O420","N421","O421"], "COMPRAS MES foto julio")

# Escanear alrededor por si hay mas celdas con texto
for sh, rng in [("STOCK 1", range(440,470)), ("Registro de STOCK 1", range(405,435))]:
    print("\n--- scan %s cols L..P ---" % sh)
    sf = wf[sh]; sv = wv[sh]
    for r in rng:
        row = []
        for col in "LMNOP":
            f = sf["%s%d" % (col, r)].value
            v = sv["%s%d" % (col, r)].value
            if f is not None or v is not None:
                row.append("%s%d: f=%r v=%r" % (col, r, f, v))
        if row:
            print("   " + " | ".join(row))
