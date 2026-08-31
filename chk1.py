# -*- coding: utf-8 -*-
import openpyxl

P = "C:/Users/usuario/Downloads/Registro KV Excel.xlsx"
wbf = openpyxl.load_workbook(P, data_only=False)
wbv = openpyxl.load_workbook(P, data_only=True)

print("=== HOJAS ===")
print(wbf.sheetnames)

def dump(sheet, cells):
    f = wbf[sheet]; v = wbv[sheet]
    for c in cells:
        print(f"  {sheet}!{c}  F={f[c].value!r}   V={v[c].value!r}")

print("\n=== STOCK 1 N/O 450-460 ===")
f = wbf['STOCK 1']; v = wbv['STOCK 1']
for r in range(448, 462):
    for col in ('M','N','O','P'):
        cf = f[f"{col}{r}"].value
        cv = v[f"{col}{r}"].value
        if cf is not None or cv is not None:
            print(f"  {col}{r}  F={cf!r}  V={cv!r}")

print("\n=== Registro de STOCK 1 N/O 414-426 ===")
f = wbf['Registro de STOCK 1']; v = wbv['Registro de STOCK 1']
for r in range(412, 428):
    for col in ('M','N','O','P'):
        cf = f[f"{col}{r}"].value
        cv = v[f"{col}{r}"].value
        if cf is not None or cv is not None:
            print(f"  {col}{r}  F={cf!r}  V={cv!r}")
