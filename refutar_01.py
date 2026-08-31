# -*- coding: utf-8 -*-
import openpyxl
from collections import OrderedDict

P = "C:/Users/usuario/Downloads/Registro KV Excel.xlsx"
wb = openpyxl.load_workbook(P, data_only=True)
wbf = openpyxl.load_workbook(P, data_only=False)

print("HOJAS:", wb.sheetnames)
print()

# 1. Verificar las formulas literales de N455:O457 en STOCK 1
s1f = wbf["STOCK 1"]
print("=== STOCK 1 formulas N/O 450-460 ===")
for r in range(450, 461):
    for c in ("N", "O", "P"):
        v = s1f[f"{c}{r}"].value
        if v not in (None, ""):
            print(f"  {c}{r} = {v!r}")

s1 = wb["STOCK 1"]
print("=== STOCK 1 valores N/O 450-460 ===")
for r in range(450, 461):
    for c in ("N", "O", "P"):
        v = s1[f"{c}{r}"].value
        if v not in (None, ""):
            print(f"  {c}{r} = {v!r}")
print()

r1f = wbf["Registro de STOCK 1"]
print("=== Registro de STOCK 1 formulas N/O 414-425 ===")
for r in range(414, 426):
    for c in ("N", "O", "P"):
        v = r1f[f"{c}{r}"].value
        if v not in (None, ""):
            print(f"  {c}{r} = {v!r}")
r1 = wb["Registro de STOCK 1"]
print("=== Registro de STOCK 1 valores N/O 414-425 ===")
for r in range(414, 426):
    for c in ("N", "O", "P"):
        v = r1[f"{c}{r}"].value
        if v not in (None, ""):
            print(f"  {c}{r} = {v!r}")
