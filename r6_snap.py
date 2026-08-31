# -*- coding: utf-8 -*-
import openpyxl, datetime, collections
P = "C:/Users/usuario/Downloads/Registro KV Excel.xlsx"
wbv = openpyxl.load_workbook(P, data_only=True)
wbf = openpyxl.load_workbook(P, data_only=False)

print("### FORMULAS de COMPRAS MES")
for sh,cells in (("STOCK 1",["N455","O455","N456","O456","N457","O457"]),
                 ("Registro de STOCK 1",["N419","O419","N420","O420","N421","O421"])):
    ws=wbf[sh]; wv=wbv[sh]
    for c in cells:
        print("   %-22s %-5s formula=%-45r valor=%r" % (sh,c,ws[c].value,wv[c].value))
print()

print("### Fechas maximas col A por hoja")
for sh in ["STOCK 1","STOCK 2","Registro de STOCK 1","Registro de STOCK 2","FLUJO","Registro de FLUJO"]:
    ws=wbv[sh]
    fs=[ws.cell(r,1).value for r in range(2,ws.max_row+1)]
    fs=[x for x in fs if isinstance(x,datetime.datetime)]
    print("   %-22s n=%4d  min=%s  max=%s" % (sh,len(fs),min(fs) if fs else None,max(fs) if fs else None))
print()

print("### Conteo de filas con CODIGO en los rangos")
def rows(sh,r0,r1):
    ws=wbv[sh]; out=[]
    for r in range(r0,r1+1):
        v=[ws.cell(r,c).value for c in range(1,9)]
        if all(x is None for x in v): continue
        out.append((r,v[0],str(v[1]).strip() if v[1] else None,v[4],v[5],v[6],v[7]))
    return out
for sh,r0,r1 in (("STOCK 1",197,454),("Registro de STOCK 1",197,418),("STOCK 2",600,695),("Registro de STOCK 2",600,695)):
    rr=rows(sh,r0,r1)
    cods=[x[2] for x in rr]
    print("   %-22s filas=%3d  con_cod=%3d  cod_unicos=%3d  sin_cod=%d" % (sh,len(rr),sum(1 for c in cods if c),len(set(c for c in cods if c)),sum(1 for c in cods if not c)))
print()

print("### Filas fechadas en AGOSTO 2026 en STOCK vivo (rangos de la formula)")
ago=[]
for sh,r0,r1 in (("STOCK 1",197,454),("STOCK 2",600,695)):
    for r,f,cod,pc,co,pf,est in rows(sh,r0,r1):
        if isinstance(f,datetime.datetime) and f.year==2026 and f.month==8:
            ago.append((sh,r,f,cod,pc,co,pf,est))
print("   n=%d  sumG=%.2f  sumE=%.2f" % (len(ago),sum(float(x[6]) for x in ago if isinstance(x[6],(int,float))),
                                          sum(float(x[4]) for x in ago if isinstance(x[4],(int,float)))))
print("   codigos agosto:", sorted(set(x[3] for x in ago)))
print()
print("### Esas mismas filas de agosto, en las hojas Registro (misma fila)")
tot=0.0; presentes=0
for sh,r,f,cod,pc,co,pf,est in ago:
    rs = "Registro de "+sh
    ws=wbv[rs]
    v=[ws.cell(r,c).value for c in range(1,9)]
    if any(x is not None for x in v):
        presentes+=1
        if isinstance(v[6],(int,float)): tot+=float(v[6])
print("   filas de agosto que YA existian en la hoja Registro (misma fila): %d de %d, sumaG en foto=%.2f" % (presentes,len(ago),tot))
print()
print("   muestra:")
for sh,r,f,cod,pc,co,pf,est in ago[:12]:
    ws=wbv["Registro de "+sh]
    v=[ws.cell(r,c).value for c in range(1,9)]
    print("     %-8s r%-4d %s %-9s  LIVE E=%-9s G=%-9s | FOTO A=%s cod=%s E=%s G=%s" %
          (sh,r,str(f)[:10],cod,pc,pf,str(v[0])[:10],v[1],v[4],v[6]))
