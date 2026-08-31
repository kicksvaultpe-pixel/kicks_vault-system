# -*- coding: utf-8 -*-
import openpyxl, datetime, collections
P = "C:/Users/usuario/Downloads/Registro KV Excel.xlsx"
wb = openpyxl.load_workbook(P, data_only=True)
def load(s):
    ws=wb[s]; rows=[]
    for r in range(2, ws.max_row+1):
        v=[ws.cell(r,c).value for c in range(1,10)]
        if all(x is None for x in v[:8]): continue
        rows.append(dict(sh=s,row=r,fecha=v[0],accion=v[1],cod=v[2],desc=v[3],talla=v[4],ing=v[5],egr=v[6],usd=v[7],info=v[8]))
    return rows
rg=load("Registro de FLUJO"); fl=load("FLUJO")
def m(v): return "%04d-%02d"%(v.year,v.month) if isinstance(v,datetime.datetime) else None
def A(r): return str(r['accion']).strip().upper() if r['accion'] else ""

print("### TODAS las filas COMPRA jun+jul 2026 en 'Registro de FLUJO' (foto 31-jul)")
sel=[r for r in rg if A(r)=="COMPRA" and m(r['fecha']) in ("2026-06","2026-07")]
sel.sort(key=lambda r:r['row'])
tot=0.0
for r in sel:
    e=r['egr']
    if e is not None: tot+=float(e)
    print("  f%-5d %-19s %-9s S/=%-10s US$=%-9s %s" % (r['row'],str(r['fecha'])[:19],r['cod'],e,r['usd'],str(r['info'])[:40]))
print("  TOTAL soles REG jun+jul COMPRA = %.2f  (n=%d)"%(tot,len(sel)))

c=collections.Counter(str(r['cod']) for r in sel)
print("\n  codigos duplicados en REG:", {k:v for k,v in c.items() if v>1})

print("\n### Filas de CUALQUIER accion en REG cuyo codigo esta en la lista de 20")
lista=['P-00961','P-00962','P-00963','P-00964','P-00967','P-00969','P-00972','P-00975','P-00976','P-00978','P-00979','P-00980','P-00981','P-00982','P-00983','P-00984','P-00986','P-00987','P-00988','P-00996']
s2=[r for r in rg if str(r['cod']).strip() in lista]
s2.sort(key=lambda r:r['row'])
tt=0.0
for r in s2:
    e=r['egr']
    if e is not None: tt+=float(e)
    print("  f%-5d %-19s %-8s %-9s ING=%-9s S/=%-10s US$=%-9s %s"%(r['row'],str(r['fecha'])[:19],A(r),r['cod'],r['ing'],e,r['usd'],str(r['info'])[:35]))
print("  suma egreso soles de esas filas en REG = %.2f (n=%d)"%(tt,len(s2)))

print("\n### Mismos codigos en FLUJO vivo, cualquier accion")
s3=[r for r in fl if str(r['cod']).strip() in lista]
s3.sort(key=lambda r:r['row'])
t3=0.0
for r in s3:
    e=r['egr']
    if e is not None: t3+=float(e)
    print("  f%-5d %-19s %-8s %-9s ING=%-9s S/=%-10s US$=%-9s %s"%(r['row'],str(r['fecha'])[:19],A(r),r['cod'],r['ing'],e,r['usd'],str(r['info'])[:35]))
print("  suma egreso soles = %.2f (n=%d)"%(t3,len(s3)))
